"""Generate Viator Cities Catalog Excel report based on the latest database records.
Matches the structure, styling, and sheets of Viator_Cities_Catalog.xlsx:
  1. All Cities (Alphabetical list of unique cities with country, status breakdowns, and accounts)
  2. Cities by Country (Grouped and sorted by country and city)
  3. Country Summary (Aggregated totals per country)
"""
import io
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def build_cities_workbook_bytes(con, account: str | None = None, allowed_accounts: list | None = None) -> bytes:
    query = """
        SELECT p.id, TRIM(p.location) AS city, p.status_canonical,
               a.id AS account_id,
               COALESCE(NULLIF(TRIM(a.name), ''), a.viator_account_id) AS account_name,
               COALESCE(NULLIF(TRIM(a.country), ''), 'Unknown') AS country,
               a.viator_account_id
        FROM products p
        JOIN accounts a ON a.id = p.account_id
        WHERE p.location IS NOT NULL AND TRIM(p.location) != ''
    """
    args = []
    if account:
        query += " AND a.viator_account_id=?"
        args.append(account)
    if allowed_accounts is not None:
        query += (f" AND a.viator_account_id IN ({','.join('?' * len(allowed_accounts))})"
                  if allowed_accounts else " AND 1=0")
        args += allowed_accounts
    query += " ORDER BY p.location"

    rows = con.execute(query, tuple(args)).fetchall()

    # Data structures for aggregation
    city_data = defaultdict(lambda: {
        'countries': set(),
        'total_products': 0,
        'live_products': 0,
        'draft_products': 0,
        'other_products': 0,
        'accounts': set(),
        'account_names': set(),
    })

    country_city_data = defaultdict(lambda: {
        'total_products': 0,
        'live_products': 0,
        'draft_products': 0,
        'accounts': set(),
        'account_names': set(),
    })

    country_summary = defaultdict(lambda: {
        'cities': set(),
        'total_products': 0,
        'live_products': 0,
        'draft_products': 0,
        'accounts': set(),
    })

    for r in rows:
        city = (r['city'] or '').strip()
        if not city:
            continue
        country = (r['country'] or 'Unknown').strip()
        status = (r['status_canonical'] or '').upper()
        acct_id = r['account_id']
        acct_name = (r['account_name'] or str(acct_id)).strip()

        # 1. City aggregate
        cd = city_data[city]
        cd['countries'].add(country)
        cd['total_products'] += 1
        if status == 'LIVE':
            cd['live_products'] += 1
        elif status == 'DRAFT':
            cd['draft_products'] += 1
        else:
            cd['other_products'] += 1
        cd['accounts'].add(acct_id)
        cd['account_names'].add(acct_name)

        # 2. Country-City aggregate
        ccd = country_city_data[(country, city)]
        ccd['total_products'] += 1
        if status == 'LIVE':
            ccd['live_products'] += 1
        elif status == 'DRAFT':
            ccd['draft_products'] += 1
        ccd['accounts'].add(acct_id)
        ccd['account_names'].add(acct_name)

        # 3. Country summary
        cs = country_summary[country]
        cs['cities'].add(city)
        cs['total_products'] += 1
        if status == 'LIVE':
            cs['live_products'] += 1
        elif status == 'DRAFT':
            cs['draft_products'] += 1
        cs['accounts'].add(acct_id)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove blank default sheet

    # Styling elements
    header_fill = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")  # Deep violet
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    total_fill = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")  # Soft violet accent
    total_font = Font(name="Segoe UI", size=11, bold=True, color="4C1D95")

    data_font = Font(name="Segoe UI", size=10, color="1E293B")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="E2E8F0")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    total_top_side = Side(border_style="thin", color="7C3AED")
    total_bottom_side = Side(border_style="double", color="7C3AED")
    total_border = Border(left=thin_border_side, right=thin_border_side, top=total_top_side, bottom=total_bottom_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # -------------------------------------------------------------------------
    # SHEET 1: All Cities (Alphabetical)
    # -------------------------------------------------------------------------
    ws1 = wb.create_sheet(title="All Cities")
    ws1.views.sheetView[0].showGridLines = True

    headers1 = [
        ("#", align_center),
        ("City / Destination", align_left),
        ("Country", align_left),
        ("Total Products", align_center),
        ("Live Products", align_center),
        ("Draft Products", align_center),
        ("Supplier Accounts", align_center),
        ("Supplier Accounts Summary", align_left)
    ]

    ws1.append([h[0] for h in headers1])
    for col_idx, (_, align) in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align
        cell.border = thin_border
    ws1.row_dimensions[1].height = 26

    sorted_cities = sorted(city_data.keys(), key=lambda c: c.lower())

    for idx, city in enumerate(sorted_cities, 1):
        cd = city_data[city]
        countries_str = ", ".join(sorted(cd['countries']))
        accounts_list = sorted(cd['account_names'])
        if len(accounts_list) <= 3:
            accounts_str = "; ".join(accounts_list)
        else:
            accounts_str = f"{'; '.join(accounts_list[:3])} (+{len(accounts_list)-3} more)"

        row_vals = [
            idx,
            city,
            countries_str,
            cd['total_products'],
            cd['live_products'],
            cd['draft_products'],
            len(cd['accounts']),
            accounts_str
        ]
        ws1.append(row_vals)
        r_idx = idx + 1
        ws1.row_dimensions[r_idx].height = 20

        is_even = (idx % 2 == 0)
        for c_idx in range(1, len(row_vals) + 1):
            cell = ws1.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill
            cell.alignment = headers1[c_idx - 1][1]

    last_r1 = len(sorted_cities) + 1
    tot_r1 = last_r1 + 1
    total_row_1 = [
        "TOTAL",
        f"{len(sorted_cities)} unique cities",
        f"{len(country_summary)} countries",
        f"=SUM(D2:D{last_r1})",
        f"=SUM(E2:E{last_r1})",
        f"=SUM(F2:F{last_r1})",
        "",
        ""
    ]
    ws1.append(total_row_1)
    ws1.row_dimensions[tot_r1].height = 24
    for c_idx in range(1, len(total_row_1) + 1):
        cell = ws1.cell(row=tot_r1, column=c_idx)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = total_border
        cell.alignment = headers1[c_idx - 1][1]

    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = f"A1:H{last_r1}"

    # -------------------------------------------------------------------------
    # SHEET 2: Cities by Country
    # -------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="Cities by Country")
    ws2.views.sheetView[0].showGridLines = True

    headers2 = [
        ("#", align_center),
        ("Country", align_left),
        ("City / Destination", align_left),
        ("Total Products", align_center),
        ("Live Products", align_center),
        ("Draft Products", align_center),
        ("Supplier Accounts", align_center),
        ("Associated Accounts", align_left)
    ]

    ws2.append([h[0] for h in headers2])
    for col_idx, (_, align) in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align
        cell.border = thin_border
    ws2.row_dimensions[1].height = 26

    sorted_country_city = sorted(country_city_data.keys(), key=lambda x: (x[0].lower(), x[1].lower()))

    for idx, (country, city) in enumerate(sorted_country_city, 1):
        ccd = country_city_data[(country, city)]
        accounts_list = sorted(ccd['account_names'])
        if len(accounts_list) <= 3:
            accounts_str = "; ".join(accounts_list)
        else:
            accounts_str = f"{'; '.join(accounts_list[:3])} (+{len(accounts_list)-3} more)"

        row_vals = [
            idx,
            country,
            city,
            ccd['total_products'],
            ccd['live_products'],
            ccd['draft_products'],
            len(ccd['accounts']),
            accounts_str
        ]
        ws2.append(row_vals)
        r_idx = idx + 1
        ws2.row_dimensions[r_idx].height = 20

        is_even = (idx % 2 == 0)
        for c_idx in range(1, len(row_vals) + 1):
            cell = ws2.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill
            cell.alignment = headers2[c_idx - 1][1]

    last_r2 = len(sorted_country_city) + 1
    tot_r2 = last_r2 + 1
    total_row_2 = [
        "TOTAL",
        f"{len(country_summary)} countries",
        f"{len(sorted_country_city)} entries",
        f"=SUM(D2:D{last_r2})",
        f"=SUM(E2:E{last_r2})",
        f"=SUM(F2:F{last_r2})",
        "",
        ""
    ]
    ws2.append(total_row_2)
    ws2.row_dimensions[tot_r2].height = 24
    for c_idx in range(1, len(total_row_2) + 1):
        cell = ws2.cell(row=tot_r2, column=c_idx)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = total_border
        cell.alignment = headers2[c_idx - 1][1]

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:H{last_r2}"

    # -------------------------------------------------------------------------
    # SHEET 3: Country Summary
    # -------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="Country Summary")
    ws3.views.sheetView[0].showGridLines = True

    headers3 = [
        ("#", align_center),
        ("Country", align_left),
        ("Total Cities", align_center),
        ("Total Products", align_center),
        ("Live Products", align_center),
        ("Draft Products", align_center),
        ("Supplier Accounts", align_center)
    ]

    ws3.append([h[0] for h in headers3])
    for col_idx, (_, align) in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align
        cell.border = thin_border
    ws3.row_dimensions[1].height = 26

    sorted_countries = sorted(country_summary.keys(), key=lambda c: (-country_summary[c]['total_products'], c.lower()))

    for idx, country in enumerate(sorted_countries, 1):
        cs = country_summary[country]
        row_vals = [
            idx,
            country,
            len(cs['cities']),
            cs['total_products'],
            cs['live_products'],
            cs['draft_products'],
            len(cs['accounts'])
        ]
        ws3.append(row_vals)
        r_idx = idx + 1
        ws3.row_dimensions[r_idx].height = 20

        is_even = (idx % 2 == 0)
        for c_idx in range(1, len(row_vals) + 1):
            cell = ws3.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill
            cell.alignment = headers3[c_idx - 1][1]

    last_r3 = len(sorted_countries) + 1
    tot_r3 = last_r3 + 1
    total_row_3 = [
        "TOTAL",
        f"{len(sorted_countries)} countries",
        f"=SUM(C2:C{last_r3})",
        f"=SUM(D2:D{last_r3})",
        f"=SUM(E2:E{last_r3})",
        f"=SUM(F2:F{last_r3})",
        f"=SUM(G2:G{last_r3})"
    ]
    ws3.append(total_row_3)
    ws3.row_dimensions[tot_r3].height = 24
    for c_idx in range(1, len(total_row_3) + 1):
        cell = ws3.cell(row=tot_r3, column=c_idx)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = total_border
        cell.alignment = headers3[c_idx - 1][1]

    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:G{last_r3}"

    # Auto-adjust column widths across all sheets
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = str(val)
                    if not val_str.startswith("="):
                        max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 55)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
